"""Damodaran XLSX parser: explicit columns; SHA-256 and cell provenance."""
from hashlib import sha256
from pathlib import Path
from decimal import Decimal,InvalidOperation
def parse_xlsx(path,*,sheet,first_row,industry_col,columns,source_url,published_at,region):
    from openpyxl import load_workbook
    if not source_url.startswith("https://pages.stern.nyu.edu/") or not published_at or not region or not columns:
        raise ValueError("Source metadata required")
    digest=sha256(Path(path).read_bytes()).hexdigest()
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        if sheet not in wb.sheetnames: raise ValueError("Unknown sheet")
        ws=wb[sheet]; result=[]
        for row in range(first_row,ws.max_row+1):
            industry=ws[f"{industry_col}{row}"].value
            if not isinstance(industry,str) or not industry.strip(): continue
            for name,col in columns.items():
                cell=f"{col}{row}"; raw=ws[cell].value
                try:
                    value=Decimal(str(raw)) if raw is not None and not isinstance(raw,bool) else None
                    if value is not None and not value.is_finite(): value=None
                except (InvalidOperation,ValueError): value=None
                result.append(dict(industry=industry.strip(),metric=name,value=str(value) if value is not None else None,region=region,source_url=source_url,source_sha256=digest,published_at=published_at,sheet=sheet,cell=cell,status="PROVISIONAL" if value is not None else "BLOCKED"))
        return result
    finally: wb.close()
