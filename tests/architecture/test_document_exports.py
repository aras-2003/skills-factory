"""Actual ARC-06 container parity. Opt in after installing export dependencies."""
import hashlib
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from scripts.architecture.document_exports import _pages, build_document_exports
from scripts.architecture.export_bundle import ExportError
from test_export_bundle import fixture


@unittest.skipUnless(os.environ.get("ARCHIFY_DOCUMENT_EXPORTS") == "1", "optional export dependencies required")
class DocumentExportTests(unittest.TestCase):
    def test_tall_export_pagination_keeps_all_rows_without_a_tiny_tail_page(self):
        from PIL import Image
        image = Image.new("RGB", (1100, 1301))
        tiles = list(_pages(image))
        self.assertEqual(sum(tile.height for tile in tiles), image.height)
        self.assertTrue(all(400 <= tile.height <= 650 for tile in tiles))

    def test_same_model_all_containers_preserve_pixels_text_provenance_and_unknowns(self):
        from docx import Document
        from pptx import Presentation
        from pypdf import PdfReader
        from PIL import Image
        from io import BytesIO
        model = fixture()
        model["nodes"][1]["label"] = "Raport źródłowy i relacje"
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "bundle"
            manifest = build_document_exports(model, out, node_executable=os.environ.get("ARCHIFY_NODE", "node"))
            equivalent = (out / "architecture.txt").read_text()
            expected = equivalent.splitlines()
            doc = Document(out / "architecture.docx")
            self.assertEqual([p.text for p in doc.paragraphs][-len(expected):], expected)
            self.assertIn(manifest["model_digest"], doc.core_properties.subject)
            for picture in doc.inline_shapes:
                self.assertIn("Full text equivalent", picture._inline.docPr.get("descr"))
            deck = Presentation(out / "architecture.pptx")
            for slide in deck.slides:
                self.assertEqual(slide.notes_slide.notes_text_frame.text, equivalent)
            pdf = PdfReader(out / "architecture.pdf")
            pdf_text = "\n".join(page.extract_text() for page in pdf.pages)
            for item in model["nodes"] + model["edges"]:
                self.assertIn(item["id"], pdf_text)
            self.assertIn("Raport źródłowy i relacje", pdf_text)
            self.assertIn("UNKNOWN", pdf_text)
            self.assertIn(manifest["model_digest"], pdf.metadata.subject)
            tree = pdf.trailer["/Root"]["/StructTreeRoot"]
            self.assertTrue(tree["/K"])
            for ref in tree["/K"]:
                figure = ref.get_object()
                self.assertEqual(figure["/S"], "/Figure")
                self.assertEqual(figure["/Alt"], equivalent)
            # Reconstruct all embedded raster tiles: exact source pixels, in
            # order, with neither dropped nor invented graph content.
            with Image.open(out / "architecture.png") as original:
                for container, prefix in (("docx", "word/media/"), ("pptx", "ppt/media/")):
                    with ZipFile(out / f"architecture.{container}") as archive:
                        tiles = [Image.open(BytesIO(archive.read(name))).convert("RGB")
                                 for name in sorted(archive.namelist()) if name.startswith(prefix)]
                    rebuilt = Image.new("RGB", original.size)
                    top = 0
                    for tile in tiles:
                        rebuilt.paste(tile, (0, top))
                        top += tile.height
                    self.assertEqual(top, original.height)
                    self.assertEqual(rebuilt.tobytes(), original.convert("RGB").tobytes())
            for name, artifact in manifest["artifacts"].items():
                self.assertEqual(hashlib.sha256((out / artifact["path"]).read_bytes()).hexdigest(), artifact["sha256"])
            for name in ("png", "docx", "pdf", "pptx"):
                self.assertTrue(manifest["formats"][name]["generated_here"])
            self.assertEqual(json.loads((out / "export-manifest.json").read_text())["model_digest"], manifest["model_digest"])

    def test_rasterizer_failure_never_claims_completed_document_exports(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "bundle"
            with self.assertRaises(ExportError):
                build_document_exports(fixture(), out, node_executable="/nonexistent/arc06-node")
            manifest = json.loads((out / "export-manifest.json").read_text())
            self.assertFalse(manifest["formats"]["pdf"]["generated_here"])
            self.assertNotIn("pdf", manifest["artifacts"])

    def test_invalid_model_never_launches_rasterizer_or_creates_output(self):
        model = fixture()
        model["edges"][0]["to"] = "missing"
        with tempfile.TemporaryDirectory() as tmp, patch("subprocess.run") as run:
            out = Path(tmp) / "bundle"
            with self.assertRaises(ValueError):
                build_document_exports(model, out)
            run.assert_not_called()
            self.assertFalse(out.exists())


if __name__ == "__main__":
    unittest.main()
