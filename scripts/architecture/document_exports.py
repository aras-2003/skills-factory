"""Offline, opt-in document consumers of the canonical ARC-06 static bundle.

Every format embeds the same PNG pixels and the complete text equivalent.
These are flattened exports, not interactive or editable architecture shapes.
"""
from __future__ import annotations

import hashlib
import json
import math
import subprocess
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from scripts.architecture.export_bundle import ExportError, build_export_bundle


def _pages(image):
    # A landscape page of diagram content at a readable fixed scale. Contiguous
    # crops retain every source pixel, including labels crossing page boundaries.
    height = math.ceil(image.height / math.ceil(image.height / 650))
    for top in range(0, image.height, height):
        yield image.crop((0, top, image.width, min(top + height, image.height)))


def _png_bytes(image):
    stream = BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


def _docx(target, images, equivalent, identity):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(11.7), Inches(8.3)
    section.top_margin = section.bottom_margin = Inches(0.6)
    section.left_margin = section.right_margin = Inches(0.6)
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Heading 1"].font.color.rgb = RGBColor.from_string("0F172A")
    doc.core_properties.title = "Architecture export"
    doc.core_properties.subject = identity
    for index, image in enumerate(images):
        if index:
            doc.add_page_break()
        doc.add_heading(f"Architecture diagram {index + 1}", level=1)
        picture = doc.add_picture(BytesIO(image), width=Inches(10.4))
        picture._inline.docPr.set("descr", "Architecture diagram. Full text equivalent follows. " + identity)
    doc.add_page_break()
    doc.add_heading("Architecture text equivalent", level=1)
    for line in equivalent.splitlines():
        doc.add_paragraph(line)
    doc.save(target)


def _pptx(target, images, equivalent, identity):
    # This is a product exporter, independent of any authoring-agent tooling.
    from pptx import Presentation
    from pptx.util import Inches, Pt
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(12), Inches(8)
    deck.core_properties.title = "Architecture export"
    deck.core_properties.subject = identity
    for index, image in enumerate(images):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(11), Inches(0.5))
        title.text_frame.paragraphs[0].text = f"Architecture diagram {index + 1}"
        title.text_frame.paragraphs[0].font.size = Pt(26)
        picture = slide.shapes.add_picture(BytesIO(image), Inches(0.5), Inches(0.9), width=Inches(11))
        picture._element.nvPicPr.cNvPr.set("descr", "Architecture diagram. Full text equivalent in notes. " + identity)
        slide.notes_slide.notes_text_frame.text = equivalent
    # Appendix text is paginated at a fixed legible size as well as available in
    # notes; no source IDs or evidence anchors are silently truncated.
    import textwrap
    lines = [wrapped for line in equivalent.splitlines() for wrapped in (textwrap.wrap(line, 100) or [""])]
    for offset in range(0, len(lines), 22):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(11), Inches(7.2))
        frame = box.text_frame
        frame.word_wrap = True
        for index, line in enumerate(lines[offset:offset + 22]):
            p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            p.text, p.font.size = line, Pt(17)
        slide.notes_slide.notes_text_frame.text = equivalent
    deck.save(target)


def _pdf(target, images, equivalent, identity):
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Paragraph, SimpleDocTemplate
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from matplotlib import get_data_path
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import ArrayObject, BooleanObject, DictionaryObject, NameObject, NumberObject, TextStringObject
    font = TTFont("ArchitectureDejaVu", str(Path(get_data_path()) / "fonts/ttf/DejaVuSans.ttf"))
    missing = {char for char in equivalent if not char.isspace() and ord(char) not in font.face.charToGlyph}
    if missing:
        raise ExportError("PDF font cannot represent every source character; export is incomplete")
    pdfmetrics.registerFont(font)
    image_pdf, text_pdf = BytesIO(), BytesIO()
    canvas = Canvas(image_pdf, pagesize=(842, 595), invariant=1)
    canvas.setTitle("Architecture export")
    for index, image in enumerate(images):
        canvas.setFont("Helvetica-Bold", 18)
        canvas.drawString(32, 559, f"Architecture diagram {index + 1}")
        # Mark each diagram as a Figure and bind its full equivalent below.
        canvas.addLiteral("/Figure <</MCID 0>> BDC")
        canvas.drawImage(ImageReader(BytesIO(image)), 32, 50, width=778, height=460,
                         preserveAspectRatio=True, anchor="nw")
        canvas.addLiteral("EMC")
        canvas.showPage()
    canvas.save()
    style = ParagraphStyle("Equivalent", fontName="ArchitectureDejaVu", fontSize=10, leading=14,
                           splitLongWords=True, spaceAfter=3)
    SimpleDocTemplate(text_pdf, pagesize=(842, 595), leftMargin=32, rightMargin=32,
                      topMargin=32, bottomMargin=32).build([
        Paragraph(escape(line) or "&#160;", style) for line in equivalent.splitlines()
    ])
    writer = PdfWriter()
    for stream in (image_pdf, text_pdf):
        writer.append(PdfReader(stream))
    tree = DictionaryObject({NameObject("/Type"): NameObject("/StructTreeRoot")})
    tree_ref = writer._add_object(tree)
    figures = ArrayObject()
    parents = ArrayObject()
    for index in range(len(images)):
        page = writer.pages[index]
        page[NameObject("/StructParents")] = NumberObject(index)
        figure = DictionaryObject({
            NameObject("/Type"): NameObject("/StructElem"), NameObject("/S"): NameObject("/Figure"),
            NameObject("/P"): tree_ref, NameObject("/Pg"): page.indirect_reference,
            NameObject("/K"): NumberObject(0), NameObject("/Alt"): TextStringObject(equivalent),
        })
        figure_ref = writer._add_object(figure)
        figures.append(figure_ref)
        parents.extend([NumberObject(index), ArrayObject([figure_ref])])
    tree[NameObject("/K")] = figures
    tree[NameObject("/ParentTree")] = writer._add_object(DictionaryObject({NameObject("/Nums"): parents}))
    writer.root_object[NameObject("/StructTreeRoot")] = tree_ref
    writer.root_object[NameObject("/MarkInfo")] = DictionaryObject({NameObject("/Marked"): BooleanObject(True)})
    writer.root_object[NameObject("/Lang")] = TextStringObject("en")
    writer.add_metadata({"/Title": "Architecture export", "/Subject": identity})
    writer.write(target)


def build_document_exports(model, output_dir, *, node_executable="node"):
    """Generate PNG/DOCX/PDF/PPTX plus the original SVG/text/manifest.

    Requires requirements-architecture-exports.txt and npm ci in this module's
    directory. No network or untrusted HTML is used during export.
    """
    from PIL import Image
    manifest = build_export_bundle(model, output_dir)
    target = Path(output_dir).expanduser().absolute()
    png = target / "architecture.png"
    try:
        result = subprocess.run([str(node_executable), str(Path(__file__).with_name("rasterize-export.cjs")),
                                 str(target / "architecture.svg"), str(png)],
                                capture_output=True, text=True, timeout=60, check=False)
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise ExportError("offline SVG rasterizer unavailable or timed out; export is incomplete") from exc
    if result.returncode or not png.is_file():
        raise ExportError("offline SVG rasterization failed; export is incomplete")
    equivalent = (target / "architecture.txt").read_text(encoding="utf-8")
    identity = f"source_revision={manifest['source_revision']}; model_digest={manifest['model_digest']}"
    with Image.open(png) as image:
        images = [_png_bytes(page) for page in _pages(image)]
    for name, builder in (("docx", _docx), ("pdf", _pdf), ("pptx", _pptx)):
        builder(target / f"architecture.{name}", images, equivalent, identity)
    manifest.pop("manifest_sha256", None)
    # Copy the nested matrix; the SVG-only builder's global contract is unchanged.
    manifest["formats"] = {key: dict(value) for key, value in manifest["formats"].items()}
    for name in ("png", "docx", "pdf", "pptx"):
        path = target / f"architecture.{name}"
        manifest["artifacts"][name] = {"path": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
        manifest["formats"][name]["generated_here"] = True
        manifest["formats"][name]["embedding"] = "same SVG-derived PNG, complete text equivalent, source revision and model digest"
    manifest["limitations"] = ["Interactive HTML remains a separate ARC-03 artifact.",
                               "PPTX diagram pixels are flattened; native editable shapes are not generated.",
                               "PDF figures have alt text; full PDF/UA conformance is not certified.",
                               "Rasterization is limited to 24 million pixels; oversized exports fail explicitly."]
    encoded = (json.dumps(manifest, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode()
    (target / "export-manifest.json").write_bytes(encoded)
    return {**manifest, "manifest_sha256": hashlib.sha256(encoded).hexdigest()}
