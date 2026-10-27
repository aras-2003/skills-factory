"""Separate executable test gate for the new Investment OS source changes."""
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "skills/investing/fundamental-intelligence"

class FundamentalIntelligenceCoverageTests(unittest.TestCase):
    def test_manifest_covers_exact_source_and_preserves_runtime_not_run(self):
        sys.path.insert(0, str(ROOT / "scripts/eval"))
        import campaign
        expected={
            "skills/investing/fundamental-intelligence/engine.py",
            "skills/investing/fundamental-intelligence/damodaran.py",
            "skills/investing/fundamental-intelligence/test_engine.py",
            "skills/investing/market-opportunity-scan/SKILL.md",
            "skills/investing/market-opportunity-scan/tests/cases.yaml",
        }
        self.assertEqual(expected,campaign.separately_validated_fundamental_intelligence_paths(ROOT))

    def test_behavior_contract_is_pinned_to_not_run_campaign(self):
        import yaml
        manifest = yaml.safe_load((ROOT / "evals/campaigns/fundamental-intelligence-2026-10/source-coverage.yaml").read_text())
        cases = yaml.safe_load((ROOT / manifest["behavior_case_suite"]).read_text())["cases"]
        selected = [case for case in cases if case["id"] in manifest["behavior_case_ids"]]
        self.assertEqual(["scan-opportunity-score-evidence-gate"], [case["id"] for case in selected])
        self.assertFalse(manifest["release_approved"])
        self.assertEqual("NOT_RUN", manifest["runtime_status"])
        self.assertTrue(all(case.get("input") and case.get("expected", {}).get("must") and case.get("expected", {}).get("must_not") for case in selected))

    def test_source_unit_suite_executes(self):
        sys.path.insert(0,str(SOURCE))
        suite=unittest.defaultTestLoader.discover(str(SOURCE),pattern="test_engine.py")
        outcome=unittest.TestResult()
        suite.run(outcome)
        self.assertGreaterEqual(outcome.testsRun,6)
        self.assertEqual([],outcome.failures)
        self.assertEqual([],outcome.errors)

    def test_damodaran_parser_cell_provenance(self):
        import tempfile
        sys.path.insert(0,str(SOURCE))
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl optional: XLSX parser integration requires dependency")
        from damodaran import parse_xlsx
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"sample.xlsx"
            wb=Workbook()
            ws=wb.active
            ws.title="Data"
            ws["A2"]="Synthetic Industry"
            ws["B2"]=3
            wb.save(path)
            rows=parse_xlsx(str(path),sheet="Data",first_row=2,industry_col="A",columns={"ev_multiple":"B"},source_url="https://pages.stern.nyu.edu/adamodar/datasets/sample.xls",published_at="2026-01-01",region="US")
            self.assertEqual(1,len(rows))
            self.assertEqual("B2",rows[0]["cell"])
            self.assertEqual("3",rows[0]["value"])
            self.assertEqual("PROVISIONAL",rows[0]["status"])

if __name__=="__main__": unittest.main()
